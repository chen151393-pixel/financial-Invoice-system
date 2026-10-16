"""七列外部群表读取；保留行号与错误，拒绝公式和重复群身份。"""

import base64
import binascii
from collections import Counter
from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from backend.core.errors import ApiError
from backend.core.xlsx_archive import check_xlsx_archive

from ..dto.group_import_dto import ExternalGroupFields

HEADERS = ["群名", "群人数", "群主", "群主所在部门", "创建时间", "外部群id", "群主id"]
CHINA = timezone(timedelta(hours=8))


def _created_at(cell):
    # 创建时间仅用于展示，缺失或无法识别不能阻止群身份导入。
    if cell.data_type in ("f", "e") or cell.value is None:
        return None
    try:
        raw = cell.value
        created = (
            raw if isinstance(raw, datetime) else datetime.fromisoformat(str(raw).strip().replace("/", "-"))
        )
        created = (created if created.tzinfo else created.replace(tzinfo=CHINA)).astimezone(UTC)
        # 超出 MySQL 日期存储范围也按缺省处理，不让展示字段阻止保存。
        return created if created.year >= 1000 else None
    except (ValueError, OverflowError):
        return None


def _member_count(cell):
    # 人数是辅助展示字段，文本或异常值不能阻止供应商群关联。
    if cell.data_type in ("f", "e") or cell.value is None:
        return None
    try:
        count = Decimal(str(cell.value).strip().removesuffix("人").strip())
        if count.is_finite() and 0 <= count <= 2147483647 and count == count.to_integral_value():
            return int(count)
    except InvalidOperation:
        pass
    return None


def _group(cells):
    if any(cell.data_type in ("f", "e") for index, cell in enumerate(cells) if index not in (1, 4)):
        raise ValueError("包含公式或 Excel 错误值，请粘贴为值")
    values = [cell.value for cell in cells]
    if len(values) != 7 or any(
        value is None or str(value).strip() == "" for index, value in enumerate(values) if index not in (1, 4)
    ):
        raise ValueError("除群人数、创建时间外，其余五列均需填写，群 ID 和群主 ID 不可缺失")
    if not isinstance(values[5], str) or not isinstance(values[6], str):
        raise ValueError("群 ID 和群主 ID 必须为文本，避免精度丢失")
    try:
        group = ExternalGroupFields(
            groupName=values[0],
            memberCount=_member_count(cells[1]),
            employee=values[2],
            department=values[3],
            createdAt=_created_at(cells[4]),
            chatId=values[5],
            userid=values[6],
        ).model_dump(mode="json")
        group["createdAt"] = group["createdAt"] or ""
        return group
    except ValidationError:
        raise ValueError("群名、群主、部门或标识格式不正确（ID 仅支持字母、数字、_@.-）") from None


def _rows(book):
    if len(book.worksheets) != 1:
        raise ApiError(422, "请只保留一张待导入工作表")
    sheet = book.worksheets[0]
    sheet.reset_dimensions()
    rows, seen_header = [], False
    for row_no, cells in enumerate(sheet.iter_rows(), 1):
        if row_no > 1001:
            raise ApiError(422, "每次最多导入 1000 行，请拆分文件")
        if row_no == 1:
            header = [str(c.value or "").strip().lower() for c in cells[:7]]
            if header != HEADERS or any(c.value is not None for c in cells[7:]):
                raise ApiError(
                    422,
                    "表头须为：群名、群人数、群主、群主所在部门、创建时间、外部群id、群主id；请补齐两列 ID",
                )
            seen_header = True
            continue
        if all(c.value is None for c in cells):
            continue
        item = dict(rowNo=row_no, groupName=str(cells[0].value or "")[:200], issue="", group=None)
        try:
            if any(c.value is not None for c in cells[7:]):
                raise ValueError("存在七列以外的数据")
            item["group"] = _group(cells[:7])
        except ValueError as error:
            item["issue"] = str(error)
        rows.append(item)
    if not seen_header or not rows:
        raise ApiError(422, "工作表没有可预览的数据")
    counts = Counter(row["group"]["chatId"] for row in rows if row["group"])
    for row in rows:
        if row["group"] and counts[row["group"]["chatId"]] > 1:
            row["issue"] = "外部群 ID 在文件中重复，请合并后重新上传"
    return rows


def read_groups(body):
    if not body.filename.lower().endswith(".xlsx"):
        raise ApiError(422, "仅支持普通 .xlsx 文件")
    try:
        raw = base64.b64decode(body.content, validate=True)
        if not raw or len(raw) > 4 * 1024 * 1024:
            raise ApiError(422, "文件不能为空且不能超过 4MB")
        check_xlsx_archive(raw)
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            return _rows(book)
        finally:
            book.close()
    except (
        binascii.Error,
        BadZipFile,
        InvalidFileException,
        ParseError,
        DefusedXmlException,
        KeyError,
        ValueError,
        OSError,
        TypeError,
    ) as error:
        raise ApiError(422, "文件无法读取，请重新导出有效的 XLSX 文件") from error
