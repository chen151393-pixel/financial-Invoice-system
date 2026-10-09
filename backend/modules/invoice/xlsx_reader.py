"""读取 XLSX 为与 XLS 相同的原始行结构，业务校验由 parser 统一执行。"""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.utils.exceptions import InvalidFileException

from backend.core.errors import ApiError


def sheet_rows(sheet, headers, keys=None):
    # 不信任工作表声明的尺寸；设置硬上限读取，避免伪造尺寸绕过限制。
    sheet.reset_dimensions()
    result = []
    for row_no, cells in enumerate(sheet.iter_rows(), 1):
        if row_no > 6001 or any(c.value is not None for c in cells[len(headers) :]):
            raise ApiError(422, f"{sheet.title}列数不符或超过6000行，请拆分文件")
        if row_no == 1:
            if [c.value for c in cells[: len(headers)]] != headers:
                raise ApiError(422, f"{sheet.title}表头与柠檬云导出模板不一致")
            continue
        values = {}
        padded = list(cells[: len(headers)]) + [EMPTY_CELL] * max(0, len(headers) - len(cells))
        for name, key, cell in zip(headers, keys or headers, padded, strict=True):
            value = cell.value
            label = f"{sheet.title}第{row_no}行{name}"
            if cell.data_type == "f":
                raise ApiError(422, f"{label}包含公式，请粘贴为值后重新导入")
            if cell.data_type == "e":
                raise ApiError(422, f"{label}为Excel错误值")
            if (
                name in {"发票代码", "发票号码", "纳税人识别号"}
                and value is not None
                and not isinstance(value, str)
            ):
                raise ApiError(422, f"{label}须为文本，避免长号码精度丢失")
            if isinstance(value, (date, datetime)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                number = Decimal(str(value))
                if name == "税率" and cell.number_format in {"0%", "0.0%", "0.00%"}:
                    value = format((number * 100).normalize(), "f") + "%"
                else:
                    value = format(number.normalize(), "f")
            else:
                value = "" if value is None else str(value).strip()
            if len(value) > 4000:
                raise ApiError(422, f"{label}过长")
            values[key] = value or None
        if any(v is not None for v in values.values()):
            result.append((row_no, values))
    return result


def read_xlsx(raw, head_headers, line_headers):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
                raise ApiError(422, "XLSX解压后过大，请拆分文件")
            if any("vbaproject" in entry.filename.lower() for entry in entries):
                raise ApiError(422, "不支持包含宏的工作簿，请导出普通XLSX文件")
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            combined = head_headers + line_headers[4:]
            candidates = [
                sheet
                for sheet in book
                if list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())) == combined
            ]
            if len(candidates) > 1:
                raise ApiError(422, "存在多个合并发票工作表，请每次保留一张待导入表")
            if candidates:
                if {"发票信息", "货物信息"} <= set(book.sheetnames):
                    raise ApiError(422, "同时存在合并表和双表模板，请只保留一种导入结构")
                sheet = candidates[0]
                keys = head_headers + ["明细_" + name for name in line_headers[4:]]
                return sheet_rows(sheet, combined, keys), None, sheet.title
            if not {"发票信息", "货物信息"} <= set(book.sheetnames):
                raise ApiError(
                    422, "未识别工作表：请使用34列票头与商品合并表，或“发票信息”“货物信息”双表模板"
                )
            return (
                sheet_rows(book["发票信息"], head_headers),
                sheet_rows(book["货物信息"], line_headers),
                None,
            )
        finally:
            book.close()
    except (
        BadZipFile,
        InvalidFileException,
        ParseError,
        DefusedXmlException,
        KeyError,
        ValueError,
        OSError,
    ) as error:
        raise ApiError(422, "无法读取XLSX文件，请重新导出有效工作簿") from error
